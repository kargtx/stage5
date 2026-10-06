import openpyxl
from openpyxl.styles import Font, Alignment

wb = openpyxl.Workbook()
ws1 = wb.active
ws1.title = "Титульный лист"

ws1.column_dimensions['A'].width = 10
ws1.column_dimensions['B'].width = 50

font_bold = Font(bold=True)
align_center = Alignment(horizontal='center', vertical='center')

c1 = ws1.cell(row=5, column=2, value="МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ")
c1.font = font_bold
c1.alignment = align_center

c2 = ws1.cell(row=7, column=2, value="Уфимский государственный авиационный технический университет")
c2.alignment = align_center

c3 = ws1.cell(row=15, column=2, value="ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ")
c3.font = Font(bold=True, size=14)
c3.alignment = align_center

c4 = ws1.cell(row=16, column=2, value="Организация параллельных потоков. Часть 1")
c4.font = Font(size=12)
c4.alignment = align_center

ws1.cell(row=25, column=2, value="Выполнил: студент гр. _______ Карим _______").alignment = Alignment(horizontal='right')
ws1.cell(row=26, column=2, value="Вариант: 4").alignment = Alignment(horizontal='right')
ws1.cell(row=28, column=2, value="Проверил: ________________").alignment = Alignment(horizontal='right')

ws1.cell(row=35, column=2, value="Уфа - 2026").alignment = align_center


ws2 = wb.create_sheet(title="Оглавление")
ws2.cell(row=1, column=1, value="Оглавление").font = font_bold
ws2.cell(row=3, column=1, value="1. Титульный лист")
ws2.cell(row=4, column=1, value="2. Оглавление")
ws2.cell(row=5, column=1, value="3. Теоретическая часть и расчеты")
ws2.cell(row=6, column=1, value="4. Результаты экспериментов")


ws3 = wb.create_sheet(title="Теоретическая часть")
ws3.column_dimensions['A'].width = 40
ws3.column_dimensions['B'].width = 25

ws3.cell(row=1, column=1, value="Функция:").font = font_bold
ws3.cell(row=1, column=2, value="0.05*x^3 + 0.3*x^2 - 20*x + 200")
ws3.cell(row=2, column=1, value="Пределы интегрирования:").font = font_bold
ws3.cell(row=2, column=2, value="A = -5, B = 20")

ws3.cell(row=4, column=1, value="Грубая оценка площади:").font = font_bold
ws3.cell(row=5, column=1, value="f(-5) =")
ws3.cell(row=5, column=2, value=301.25)
ws3.cell(row=6, column=1, value="f(20) =")
ws3.cell(row=6, column=2, value=320)
ws3.cell(row=7, column=1, value="Среднее f(x) (примерно) =")
ws3.cell(row=7, column=2, value=310)
ws3.cell(row=8, column=1, value="Площадь (S) =")
ws3.cell(row=8, column=2, value=7750)

ws3.cell(row=10, column=1, value="Аналитическое (точное) решение:").font = font_bold
ws3.cell(row=11, column=1, value="F(20) =")
ws3.cell(row=11, column=2, value=2800)
ws3.cell(row=12, column=1, value="F(-5) =")
ws3.cell(row=12, column=2, value=-1254.6875)
ws3.cell(row=13, column=1, value="Точная площадь (S) =")
ws3.cell(row=13, column=2, value=4054.6875)

ws4 = wb.create_sheet(title="Результаты")
ws4.cell(row=1, column=1, value="Здесь можно вставить результаты из CSV файлов").font = font_bold

wb.save("Отчет_ВАС_Вариант4.xlsx")
