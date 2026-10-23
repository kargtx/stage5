import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_full_practice_2(student_name="Студент Гайнанов Д.И.", filename_out="Практика_2_Решение.xlsx"):
    wb = openpyxl.Workbook()
    
    FONT_NAME = 'Segoe UI'
    
    # Fonts
    font_main_title = Font(name=FONT_NAME, size=15, bold=True, color='FFFFFF')
    font_subtitle = Font(name=FONT_NAME, size=11, bold=False, italic=True, color='FFFFFF')
    font_meta = Font(name=FONT_NAME, size=10, bold=False, color='E0E0E0')
    font_sec_header = Font(name=FONT_NAME, size=12, bold=True, color='1B365D')
    font_card_title = Font(name=FONT_NAME, size=11, bold=True, color='1B365D')
    font_tbl_header = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
    font_regular = Font(name=FONT_NAME, size=10, bold=False, color='262626')
    font_bold = Font(name=FONT_NAME, size=10, bold=True, color='262626')
    font_italic = Font(name=FONT_NAME, size=9, italic=True, color='595959')
    font_math = Font(name=FONT_NAME, size=10, italic=True, color='1F4E78')
    font_result = Font(name=FONT_NAME, size=11, bold=True, color='1E4620')
    font_status_ok = Font(name=FONT_NAME, size=10, bold=True, color='1E4620')
    
    # Fills
    fill_navy_header = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    fill_sub_header = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    fill_tbl_header = PatternFill(start_color='335E99', end_color='335E99', fill_type='solid')
    fill_card_header = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    fill_condition = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid')
    fill_zebra = PatternFill(start_color='F9FAFC', end_color='F9FAFC', fill_type='solid')
    fill_result = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_accent_row = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    
    # Borders
    border_thin_gray = Side(style='thin', color='D9D9D9')
    border_double_bottom = Side(style='double', color='1B365D')
    border_green_med = Side(style='medium', color='385723')
    border_green_thin = Side(style='thin', color='70AD47')
    
    box_cell_border = Border(left=border_thin_gray, right=border_thin_gray, top=border_thin_gray, bottom=border_thin_gray)
    box_total_border = Border(left=border_thin_gray, right=border_thin_gray, top=border_thin_gray, bottom=border_double_bottom)
    box_result_border = Border(left=border_green_thin, right=border_green_thin, top=border_green_thin, bottom=border_green_thin)
    
    # Alignments
    align_center = Alignment(horizontal='center', vertical='center')
    align_center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')
    align_left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')
    
    def apply_row_style(ws, row_idx, start_col, end_col, font=None, fill=None, border=None, alignment=None, number_format=None):
        for col_idx in range(start_col, end_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if font: cell.font = font
            if fill: cell.fill = fill
            if border: cell.border = border
            if alignment: cell.alignment = alignment
            if number_format: cell.number_format = number_format

    def make_title_banner(ws, title_text, subtitle_text, meta_text, max_col=7):
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=max_col)
        c2 = ws.cell(row=2, column=2, value=title_text)
        c2.font = font_main_title
        c2.alignment = align_center
        apply_row_style(ws, 2, 2, max_col, fill=fill_navy_header)
        ws.row_dimensions[2].height = 30

        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=max_col)
        c3 = ws.cell(row=3, column=2, value=subtitle_text)
        c3.font = font_subtitle
        c3.alignment = align_center
        apply_row_style(ws, 3, 2, max_col, fill=fill_sub_header)
        ws.row_dimensions[3].height = 24

        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=max_col)
        c4 = ws.cell(row=4, column=2, value=meta_text)
        c4.font = font_meta
        c4.alignment = align_center
        apply_row_style(ws, 4, 2, max_col, fill=fill_sub_header)
        ws.row_dimensions[4].height = 20

    RESULTS_REF = {}

    calc_col_headers = [
        "Определяемый показатель", "Обозначение", "Математическая формула", 
        "Значение (формула Excel)", "Ед. изм.", "Пояснение к результату"
    ]
    
    def render_dynamic_task(ws, task_title, task_condition, inputs, calculations, final_note, start_r, task_key):
        r = start_r
        task_start_row = r
        
        # 1. Header Card
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c_title = ws.cell(row=r, column=2, value=task_title)
        c_title.font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
        c_title.alignment = align_left
        apply_row_style(ws, r, 2, 7, fill=fill_navy_header)
        ws.row_dimensions[r].height = 24
        
        # 2. Condition Box
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=7)
        c_cond = ws.cell(row=r, column=2, value=f"Условие: {task_condition}")
        c_cond.font = font_italic
        c_cond.alignment = align_left_wrap
        apply_row_style(ws, r, 2, 7, fill=fill_condition)
        apply_row_style(ws, r+1, 2, 7, fill=fill_condition)
        ws.row_dimensions[r].height = 18
        ws.row_dimensions[r+1].height = 18
        r += 1
        
        # 3. Inputs Header
        r += 1
        ws.cell(row=r, column=2, value="Таблица 1. Исходные данные задачи").font = font_card_title
        ws.row_dimensions[r].height = 20
        
        r += 1
        sub_headers_inp = ["Наименование параметра", "Обозначение", "Значение", "Ед. изм.", "Пояснение"]
        for ci, sh in enumerate(sub_headers_inp, start=2):
            c = ws.cell(row=r, column=ci, value=sh)
            c.font = font_tbl_header
            c.fill = fill_sub_header
            c.alignment = align_center
            c.border = box_cell_border
        ws.row_dimensions[r].height = 22
        
        inp_map = {}
        for inp_row in inputs:
            r += 1
            ws.row_dimensions[r].height = 20
            p_name, p_sym, p_val, p_unit, p_desc = inp_row
            
            if isinstance(p_val, str) and '{' in p_val:
                p_val = p_val.format(**inp_map)
            
            c1 = ws.cell(row=r, column=2, value=p_name)
            c1.font = font_regular
            c1.alignment = align_left
            c1.border = box_cell_border
            
            c2 = ws.cell(row=r, column=3, value=p_sym)
            c2.font = font_bold
            c2.alignment = align_center
            c2.border = box_cell_border
            
            c3 = ws.cell(row=r, column=4, value=p_val)
            c3.font = font_bold
            c3.alignment = align_right
            c3.border = box_cell_border
            if isinstance(p_val, (int, float)):
                if abs(p_val) < 0.0001 and p_val != 0:
                    c3.number_format = '0.00E+00'
                elif isinstance(p_val, float):
                    c3.number_format = '#,##0.####'
                else:
                    c3.number_format = '#,##0'
            
            c4 = ws.cell(row=r, column=5, value=p_unit)
            c4.font = font_regular
            c4.alignment = align_center
            c4.border = box_cell_border
            
            c5 = ws.cell(row=r, column=6, value=p_desc)
            c5.font = font_italic
            c5.alignment = align_left
            c5.border = box_cell_border
            
            inp_map[p_sym] = f"D{r}"
            clean_sym = p_sym.replace('(', '_').replace(')', '').replace(' ', '').replace('^', '').replace('-', '_').replace('.', '_')
            inp_map[clean_sym] = f"D{r}"
        
        # 4. Calculations Header
        r += 2
        ws.cell(row=r, column=2, value="Таблица 2. Расчет показателей надежности").font = font_card_title
        ws.row_dimensions[r].height = 20
        
        r += 1
        for ci, ch in enumerate(calc_col_headers, start=2):
            c = ws.cell(row=r, column=ci, value=ch)
            c.font = font_tbl_header
            c.fill = fill_tbl_header
            c.alignment = align_center
            c.border = box_cell_border
        ws.row_dimensions[r].height = 24
        
        calc_map = {}
        for calc_row in calculations:
            r += 1
            ws.row_dimensions[r].height = 22
            c_name, c_sym, c_math, c_fmt_template, c_unit, c_note, c_num_fmt = calc_row
            
            resolved_formula = c_fmt_template.format(**inp_map, **calc_map)
            
            c1 = ws.cell(row=r, column=2, value=c_name)
            c1.font = font_regular
            c1.alignment = align_left
            c1.border = box_cell_border
            
            c2 = ws.cell(row=r, column=3, value=c_sym)
            c2.font = font_bold
            c2.alignment = align_center
            c2.border = box_cell_border
            
            c3 = ws.cell(row=r, column=4, value=c_math)
            c3.font = font_math
            c3.alignment = align_left
            c3.border = box_cell_border
            
            c4 = ws.cell(row=r, column=5, value=resolved_formula)
            c4.font = font_result
            c4.alignment = align_right
            c4.fill = fill_result
            c4.border = box_result_border
            c4.number_format = c_num_fmt
            
            calc_map[c_sym] = f"E{r}"
            clean_csym = c_sym.replace('(', '_').replace(')', '').replace(' ', '').replace('^', '').replace('-', '_').replace('.', '_')
            calc_map[clean_csym] = f"E{r}"
            RESULTS_REF[f"{task_key}_{clean_csym}"] = f"'{ws.title}'!E{r}"
            
            c5 = ws.cell(row=r, column=6, value=c_unit)
            c5.font = font_regular
            c5.alignment = align_center
            c5.border = box_cell_border
            
            c6 = ws.cell(row=r, column=7, value=c_note)
            c6.font = font_italic
            c6.alignment = align_left
            c6.border = box_cell_border
            
        # 5. Final Output / Answer Box
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c_ans = ws.cell(row=r, column=2, value=f"ВЫВОД / ОТВЕТ: {final_note}")
        c_ans.font = font_bold
        c_ans.alignment = align_left
        apply_row_style(ws, r, 2, 7, fill=fill_accent_row, border=Border(top=border_green_med, bottom=border_green_med))
        ws.row_dimensions[r].height = 24
        
        RESULTS_REF[f"{task_key}_start"] = f"'{ws.title}'!A{task_start_row}"
        
        r += 3
        return r

    # =========================================================================
    # CREATE SHEET: ЗАДАЧИ 6-12 (САМОСТОЯТЕЛЬНЫЕ)
    # =========================================================================
    ws_main = wb.active
    ws_main.title = "Задачи 6-12"
    ws_main.views.sheetView[0].showGridLines = True
    
    make_title_banner(
        ws_main,
        "ЗАДАЧИ ДЛЯ САМОСТОЯТЕЛЬНОГО РЕШЕНИЯ (№ 6 – 12)",
        "Практическая работа № 2. Аналитическое определение количественных характеристик надежности изделия",
        "Все расчеты выполнены формулами Excel со ссылками на ячейки исходных данных",
        max_col=7
    )
    
    cur_r = 6
    
    # --- ЗАДАЧА 6 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 6. Расчет показателей надежности системы при L1 = 460 ч (t = 800 ч)",
        "Средняя наработка до отказа L1 системы равна 460 часов. Для экспоненциального закона надежности определить вероятность безотказной работы, частоту отказов и интенсивность отказов за 800 часов работы системы.",
        [
            ("Средняя наработка до отказа", "L1", 460, "ч", "L1 = mt = 460 ч"),
            ("Время работы (наработка)", "t", 800, "ч", "Момент времени t = 800 ч")
        ],
        [
            ("Интенсивность отказов", "l_800", "1 / L1", "=1/{L1}", "1/ч", "λ = const = 1 / 460 ≈ 0.002174 ч⁻¹", "0.000000"),
            ("Показатель экспоненты (λ · t)", "lt", "λ · t", "={l_800}*{t}", "—", "λ · t = 800 / 460 ≈ 1.73913", "0.0000"),
            ("Вероятность безотказной работы", "P_800", "EXP(-λ · t)", "=EXP(-{lt})", "—", "P(800) = exp(-1.73913) ≈ 0.1757 (17.57%)", "0.0000"),
            ("Вероятность отказа (доп.)", "q_800", "1 - P(t)", "=1-{P_800}", "—", "q(800) = 1 - 0.1757 = 0.8243 (82.43%)", "0.0000"),
            ("Частота отказов системы", "f_800", "λ · P(t)", "={l_800}*{P_800}", "1/ч", "f(800) = 0.002174 · 0.1757 ≈ 0.000382 ч⁻¹", "0.000000")
        ],
        "При t = 800 ч: интенсивность отказов λ = 0,002174 1/ч (2,17·10⁻³ ч⁻¹); вероятность безотказной работы P(800) = 0,1757 (17,57%); частота отказов f(800) = 0,000382 1/ч (3,82·10⁻⁴ ч⁻¹).",
        cur_r,
        "t6"
    )

    # --- ЗАДАЧА 7 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 7. Надежность системы из 6000 последовательных элементов (t = 50 ч)",
        "Система управления состоит из 6000 элементов, соединенных последовательно, средняя интенсивность отказов которых λср(t) = 0,2∙10⁻⁶ ч⁻¹. Необходимо определить вероятность безотказной работы в течение Т = 50 ч и среднее время безотказной работы системы.",
        [
            ("Число элементов в системе", "n", 6000, "шт.", "Количество последовательных элементов"),
            ("Средняя интенсивность отказов элемента", "l_elem", 0.0000002, "1/ч", "λср = 0.2·10⁻⁶ ч⁻¹ = 2·10⁻⁷ ч⁻¹"),
            ("Время непрерывной работы", "T", 50, "ч", "Период испытаний T = 50 ч")
        ],
        [
            ("Интенсивность отказов системы", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 6000 · 0.2·10⁻⁶ = 0.0012 ч⁻¹ (1.2·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λc · T)", "lcT", "λc · T", "={lc}*{T}", "—", "λc · T = 0.0012 · 50 = 0.0600", "0.0000"),
            ("Вероятность безотказной работы системы", "Pc_50", "EXP(-λc · T)", "=EXP(-{lcT})", "—", "Pc(50) = exp(-0.06) ≈ 0.9418 (94.18%)", "0.0000"),
            ("Вероятность отказа системы (доп.)", "qc_50", "1 - Pc(T)", "=1-{Pc_50}", "—", "qc(50) = 1 - 0.9418 = 0.0582 (5.82%)", "0.0000"),
            ("Частота отказов системы (доп.)", "fc_50", "λc · Pc(T)", "={lc}*{Pc_50}", "1/ч", "fc(50) = 0.0012 · 0.9418 ≈ 0.001130 ч⁻¹", "0.000000"),
            ("Среднее время безотказной работы системы", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.0012 = 833.33 ч", "0.00")
        ],
        "Для системы из 6000 элементов: вероятность безотказной работы за 50 ч равна Pc(50) = 0,9418 (94,18%); среднее время безотказной работы системы mtc = 833,33 ч (интенсивность отказов λc = 1,2·10⁻³ ч⁻¹).",
        cur_r,
        "t7"
    )

    # --- ЗАДАЧА 8 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 8. Надежность системы из 5000 последовательных элементов (t = 120 ч)",
        "Система управления состоит из 5000 элементов, соединенных последовательно, средняя интенсивность отказов которых λср(t) = 0,3∙10⁻⁶ ч⁻¹. Необходимо определить вероятность безотказной работы в течение Т = 120 ч и среднее время безотказной работы системы.",
        [
            ("Число элементов в системе", "n", 5000, "шт.", "Количество последовательных элементов"),
            ("Средняя интенсивность отказов элемента", "l_elem", 0.0000003, "1/ч", "λср = 0.3·10⁻⁶ ч⁻¹ = 3·10⁻⁷ ч⁻¹"),
            ("Время непрерывной работы", "T", 120, "ч", "Период испытаний T = 120 ч")
        ],
        [
            ("Интенсивность отказов системы", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 5000 · 0.3·10⁻⁶ = 0.0015 ч⁻¹ (1.5·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λc · T)", "lcT", "λc · T", "={lc}*{T}", "—", "λc · T = 0.0015 · 120 = 0.1800", "0.0000"),
            ("Вероятность безотказной работы системы", "Pc_120", "EXP(-λc · T)", "=EXP(-{lcT})", "—", "Pc(120) = exp(-0.18) ≈ 0.8353 (83.53%)", "0.0000"),
            ("Вероятность отказа системы (доп.)", "qc_120", "1 - Pc(T)", "=1-{Pc_120}", "—", "qc(120) = 1 - 0.8353 = 0.1647 (16.47%)", "0.0000"),
            ("Частота отказов системы (доп.)", "fc_120", "λc · Pc(T)", "={lc}*{Pc_120}", "1/ч", "fc(120) = 0.0015 · 0.8353 ≈ 0.001253 ч⁻¹", "0.000000"),
            ("Среднее время безотказной работы системы", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.0015 = 666.67 ч", "0.00")
        ],
        "Для системы из 5000 элементов: вероятность безотказной работы за 120 ч равна Pc(120) = 0,8353 (83,53%); среднее время безотказной работы mtc = 666,67 ч (интенсивность отказов λc = 1,5·10⁻³ ч⁻¹).",
        cur_r,
        "t8"
    )

    # --- ЗАДАЧА 9 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 9. Надежность невосстанавливаемой ЭВМ из 200000 элементов в течение суток",
        "Невосстанавливаемая в процессе работы электронная машина состоит из 200000 элементов, соединенных последовательно, средняя интенсивность отказов которых λср(t) = 0,3∙10⁻⁶ ч⁻¹. Требуется определить вероятность безотказной работы электронной машины в течение суток и среднее время безотказной работы электронной машины.",
        [
            ("Число элементов машины", "n", 200000, "шт.", "Количество последовательных элементов"),
            ("Интенсивность отказов одного элемента", "l_elem", 0.0000003, "1/ч", "λср = 0.3·10⁻⁶ ч⁻¹"),
            ("Время работы (сутки)", "t", 24, "ч", "Суточный интервал работы (1 сутки = 24 ч)")
        ],
        [
            ("Интенсивность отказов машины", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 200000 · 0.3·10⁻⁶ = 0.06 ч⁻¹ (6·10⁻² ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λc · t)", "lct", "λc · t", "={lc}*{t}", "—", "λc · t = 0.06 · 24 = 1.4400", "0.0000"),
            ("Вероятность безотказной работы за сутки", "Pc_24", "EXP(-λc · t)", "=EXP(-{lct})", "—", "Pc(24) = exp(-1.44) ≈ 0.2369 (23.69%)", "0.0000"),
            ("Вероятность отказа машины за сутки", "qc_24", "1 - Pc(24)", "=1-{Pc_24}", "—", "qc(24) = 1 - 0.2369 = 0.7631 (76.31%)", "0.0000"),
            ("Частота отказов машины (доп.)", "fc_24", "λc · Pc(24)", "={lc}*{Pc_24}", "1/ч", "fc(24) = 0.06 · 0.2369 ≈ 0.014216 ч⁻¹", "0.000000"),
            ("Среднее время безотказной работы машины", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.06 = 16.67 ч (16 ч 40 мин)", "0.00")
        ],
        "Для ЭВМ из 200000 элементов: суточная вероятность безотказной работы Pc(24) = 0,2369 (23,69%); среднее время безотказной работы mtc = 16,67 ч (16 ч 40 мин; интенсивность отказов λc = 0,06 ч⁻¹).",
        cur_r,
        "t9"
    )

    # --- ЗАДАЧА 10 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 10. Расчет Pc(t), qc(t), fc(t), mtc для системы из 12600 элементов (t = 50 ч)",
        "Система состоит из 12600 элементов, средняя интенсивность отказов которых λср = 0,32∙10⁻⁶ ч⁻¹. Требуется определить Pc(t), qc(t), fc(t), mtc для t = 50 ч.",
        [
            ("Число элементов в системе", "n", 12600, "шт.", "Количество последовательных элементов"),
            ("Интенсивность отказов элемента", "l_elem", 0.00000032, "1/ч", "λср = 0.32·10⁻⁶ ч⁻¹"),
            ("Время наработки", "t", 50, "ч", "Контрольная точка t = 50 ч")
        ],
        [
            ("Интенсивность отказов системы", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 12600 · 0.32·10⁻⁶ = 0.004032 ч⁻¹ (4.032·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λc · t)", "lct", "λc · t", "={lc}*{t}", "—", "λc · t = 0.004032 · 50 = 0.2016", "0.0000"),
            ("Вероятность безотказной работы системы", "Pc_50", "EXP(-λc · t)", "=EXP(-{lct})", "—", "Pc(50) = exp(-0.2016) ≈ 0.8174 (81.74%)", "0.0000"),
            ("Вероятность отказа системы", "qc_50", "1 - Pc(t)", "=1-{Pc_50}", "—", "qc(50) = 1 - 0.8174 = 0.1826 (18.26%)", "0.0000"),
            ("Частота отказов системы", "fc_50", "λc · Pc(t)", "={lc}*{Pc_50}", "1/ч", "fc(50) = 0.004032 · 0.8174 ≈ 0.003296 ч⁻¹", "0.000000"),
            ("Среднее время безотказной работы системы", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.004032 ≈ 248.02 ч", "0.00")
        ],
        "При t = 50 ч: интенсивность отказов λc = 0,004032 1/ч; Pc(50) = 0,8174 (81,74%); qc(50) = 0,1826 (18,26%); fc(50) = 0,003296 1/ч (3,30·10⁻³ ч⁻¹); mtc = 248,02 ч.",
        cur_r,
        "t10"
    )

    # --- ЗАДАЧА 11 ---
    cur_r = render_dynamic_task(
        ws_main,
        "ЗАДАЧА № 11. Определение p(t), q(t), f(t), mt для прибора с λ = 0,1∙10⁻⁶ ч⁻¹ (t = 180 ч)",
        "Время работы прибора до отказа подчинено экспоненциальному закону распределения с параметром λ(t) = 0,1∙10⁻⁶ ч⁻¹; t = 180 ч. Требуется определить количественные характеристики надежности элемента p(t), q(t), f(t), mt.",
        [
            ("Параметр распределения (интенсивность)", "l_val", 0.0000001, "1/ч", "λ = const = 0.1·10⁻⁶ ч⁻¹ = 1·10⁻⁷ ч⁻¹"),
            ("Время наработки", "t", 180, "ч", "Момент времени t = 180 ч")
        ],
        [
            ("Показатель экспоненты (λ · t)", "lt", "λ · t", "={l_val}*{t}", "—", "λ · t = 0.1·10⁻⁶ · 180 = 0.000018", "0.000000"),
            ("Вероятность безотказной работы", "p_180", "EXP(-λ · t)", "=EXP(-{lt})", "—", "p(180) = exp(-0.000018) ≈ 0.999982 (99.9982%)", "0.000000"),
            ("Вероятность отказа прибора", "q_180", "1 - p(t)", "=1-{p_180}", "—", "q(180) = 1 - 0.999982 = 0.000018 (1.80·10⁻⁵)", "0.000000"),
            ("Частота отказов прибора", "f_180", "λ · p(t)", "={l_val}*{p_180}", "1/ч", "f(180) = 10⁻⁷ · 0.999982 ≈ 0.0000001 ч⁻¹ (1.00·10⁻⁷ ч⁻¹)", "0.00E+00"),
            ("Среднее время безотказной работы", "mt", "1 / λ", "=1/{l_val}", "ч", "mt = 1 / 10⁻⁷ = 10 000 000 ч (10 млн часов)", "#,##0")
        ],
        "Количественные характеристики прибора: p(180) = 0,999982; q(180) = 0,000018 (1,8·10⁻⁵); f(180) = 1,00·10⁻⁷ 1/ч; средняя наработка mt = 10 000 000 ч (10⁷ ч).",
        cur_r,
        "t11"
    )

    # --- ЗАДАЧА 12 (СИСТЕМА ИЗ 5 ПРИБОРОВ) ---
    r = cur_r
    task_start_row = r
    
    ws_main.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c_title12 = ws_main.cell(row=r, column=2, value="ЗАДАЧА № 12. Расчет среднего времени безотказной работы системы из пяти приборов")
    c_title12.font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
    apply_row_style(ws_main, r, 2, 7, fill=fill_navy_header)
    ws_main.row_dimensions[r].height = 24
    
    r += 1
    ws_main.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=7)
    c_cond12 = ws_main.cell(row=r, column=2, value="Условие: Система состоит из пяти приборов, среднее время безотказной работы которых равно: mt1 = 83 ч; mt2 = 220 ч; mt3 = 280 ч; mt4 = 400 ч; mt5 = 700 ч. Для приборов справедлив экспоненциальный закон надежности. Требуется найти среднее время безотказной работы системы mtc.")
    c_cond12.font = font_italic
    c_cond12.alignment = align_left_wrap
    apply_row_style(ws_main, r, 2, 7, fill=fill_condition)
    apply_row_style(ws_main, r+1, 2, 7, fill=fill_condition)
    ws_main.row_dimensions[r].height = 18
    ws_main.row_dimensions[r+1].height = 18
    r += 1
    
    # Table 1: Приборы
    r += 1
    ws_main.cell(row=r, column=2, value="Таблица 1. Исходные параметры надежности приборов системы").font = font_card_title
    ws_main.row_dimensions[r].height = 20
    
    r += 1
    t12_headers = ["№ прибора", "Обозначение", "Средняя наработка mti", "Ед. изм.", "Интенсивность отказов λi = 1/mti", "Пояснение"]
    for ci, th in enumerate(t12_headers, start=2):
        cell = ws_main.cell(row=r, column=ci, value=th)
        cell.font = font_tbl_header
        cell.fill = fill_sub_header
        cell.alignment = align_center_wrap
        cell.border = box_cell_border
    ws_main.row_dimensions[r].height = 24
    
    devices = [
        (1, "mt1", 83, "Прибор 1 (наименее надежный)"),
        (2, "mt2", 220, "Прибор 2"),
        (3, "mt3", 280, "Прибор 3"),
        (4, "mt4", 400, "Прибор 4"),
        (5, "mt5", 700, "Прибор 5 (наиболее надежный)")
    ]
    
    st_t12 = r + 1
    for dev in devices:
        r += 1
        ws_main.row_dimensions[r].height = 20
        d_idx, d_sym, d_mt, d_desc = dev
        
        c1 = ws_main.cell(row=r, column=2, value=f"Прибор {d_idx}")
        c1.font = font_regular; c1.alignment = align_center; c1.border = box_cell_border
        
        c2 = ws_main.cell(row=r, column=3, value=d_sym)
        c2.font = font_bold; c2.alignment = align_center; c2.border = box_cell_border
        
        c3 = ws_main.cell(row=r, column=4, value=d_mt)
        c3.font = font_bold; c3.alignment = align_right; c3.border = box_cell_border
        c3.number_format = '#,##0'
        
        c4 = ws_main.cell(row=r, column=5, value="ч")
        c4.font = font_regular; c4.alignment = align_center; c4.border = box_cell_border
        
        c5 = ws_main.cell(row=r, column=6, value=f"=1/D{r}")
        c5.font = font_bold; c5.alignment = align_right; c5.border = box_cell_border
        c5.number_format = '0.000000'
        
        c6 = ws_main.cell(row=r, column=7, value=d_desc)
        c6.font = font_italic; c6.alignment = align_left; c6.border = box_cell_border
    end_t12 = r
    
    # Total row
    r += 1
    ws_main.row_dimensions[r].height = 22
    ws_main.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws_main.cell(row=r, column=2, value="ИТОГО: Суммарная интенсивность отказов системы λc = Σ λi :").font = font_bold
    ws_main.cell(row=r, column=2).alignment = align_right
    apply_row_style(ws_main, r, 2, 5, fill=fill_card_header, border=box_total_border)
    
    c_sum_lc = ws_main.cell(row=r, column=6, value=f"=SUM(F{st_t12}:F{end_t12})")
    c_sum_lc.font = font_result; c_sum_lc.alignment = align_right; c_sum_lc.fill = fill_card_header; c_sum_lc.border = box_total_border
    c_sum_lc.number_format = '0.000000'
    ref_lc_t12 = f"F{r}"
    RESULTS_REF["t12_lc"] = f"'Задачи 6-12'!{ref_lc_t12}"
    
    c_sum_u = ws_main.cell(row=r, column=7, value="1/ч (λc ≈ 0.024094 ч⁻¹)")
    c_sum_u.font = font_bold; c_sum_u.alignment = align_left; c_sum_u.fill = fill_card_header; c_sum_u.border = box_total_border
    
    # Table 2: Итоговый расчет mtc
    r += 2
    ws_main.cell(row=r, column=2, value="Таблица 2. Расчет показателей надежности системы").font = font_card_title
    ws_main.row_dimensions[r].height = 20
    
    r += 1
    for ci, ch in enumerate(calc_col_headers, start=2):
        c = ws_main.cell(row=r, column=ci, value=ch)
        c.font = font_tbl_header; c.fill = fill_tbl_header; c.alignment = align_center; c.border = box_cell_border
    ws_main.row_dimensions[r].height = 24
    
    t12_calcs = [
        ("Суммарная интенсивность отказов", "lc", f"={ref_lc_t12}", f"={ref_lc_t12}", "1/ч", "λc = 1/83 + 1/220 + 1/280 + 1/400 + 1/700 ≈ 0.024094 ч⁻¹", "0.000000"),
        ("Среднее время безотказной работы системы", "mtc", f"1 / λc", f"=1/{ref_lc_t12}", "ч", "mtc = 1 / 0.024094 ≈ 41.50 ч", "0.00")
    ]
    
    for crow in t12_calcs:
        r += 1
        ws_main.row_dimensions[r].height = 22
        c_name, c_sym, c_math, c_form, c_unit, c_note, c_fmt = crow
        
        c1 = ws_main.cell(row=r, column=2, value=c_name); c1.font = font_regular; c1.alignment = align_left; c1.border = box_cell_border
        c2 = ws_main.cell(row=r, column=3, value=c_sym); c2.font = font_bold; c2.alignment = align_center; c2.border = box_cell_border
        c3 = ws_main.cell(row=r, column=4, value=c_math); c3.font = font_math; c3.alignment = align_left; c3.border = box_cell_border
        c4 = ws_main.cell(row=r, column=5, value=c_form); c4.font = font_result; c4.alignment = align_right; c4.fill = fill_result; c4.border = box_result_border; c4.number_format = c_fmt
        c5 = ws_main.cell(row=r, column=6, value=c_unit); c5.font = font_regular; c5.alignment = align_center; c5.border = box_cell_border
        c6 = ws_main.cell(row=r, column=7, value=c_note); c6.font = font_italic; c6.alignment = align_left; c6.border = box_cell_border
        
        if c_sym == "mtc":
            RESULTS_REF["t12_mtc"] = f"'Задачи 6-12'!E{r}"
            
    r += 1
    ws_main.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c_ans12 = ws_main.cell(row=r, column=2, value="ВЫВОД / ОТВЕТ: Интенсивность отказов системы λc = 0,024094 1/ч; среднее время безотказной работы системы mtc = 41,50 ч.")
    c_ans12.font = font_bold; c_ans12.alignment = align_left
    apply_row_style(ws_main, r, 2, 7, fill=fill_accent_row, border=Border(top=border_green_med, bottom=border_green_med))
    ws_main.row_dimensions[r].height = 24
    
    RESULTS_REF["t12_start"] = f"'Задачи 6-12'!A{task_start_row}"


    # =========================================================================
    # CREATE SHEET: ЗАДАЧИ 1-5 (ТИПОВЫЕ ПРИМЕРЫ ИЗ ПРЕЗЕНТАЦИИ)
    # =========================================================================
    ws_typ = wb.create_sheet(title="Задачи 1-5")
    ws_typ.views.sheetView[0].showGridLines = True
    
    make_title_banner(
        ws_typ,
        "ТИПОВЫЕ ЗАДАЧИ С РЕШЕНИЕМ ИЗ ПРЕЗЕНТАЦИИ (№ 1 – 5)",
        "Практическая работа № 2. Аналитическое определение количественных характеристик надежности изделия",
        "Полное соответствие слайдам 8–10 презентации лекций",
        max_col=7
    )
    
    cur_r_typ = 6
    
    # --- ТИПОВАЯ ЗАДАЧА 1 ---
    cur_r_typ = render_dynamic_task(
        ws_typ,
        "ЗАДАЧА № 1. Автоматический регулятор (L1 = 640 ч, t = 120 ч)",
        "Средняя наработка до отказа L1 автоматического регулятора равна 640 часов. Предполагается, что справедлив экспоненциальный закон надежности. Определить вероятность безотказной работы, частоту отказов и интенсивность отказов за 120 часов работы системы.",
        [
            ("Средняя наработка до отказа", "L1", 640, "ч", "L1 = 640 ч"),
            ("Время наработки", "t", 120, "ч", "Момент времени t = 120 ч")
        ],
        [
            ("Интенсивность отказов", "l_120", "1 / L1", "=1/{L1}", "1/ч", "λ = const = 1 / 640 ≈ 0.001563 ч⁻¹ (1.56·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λ · t)", "lt", "λ · t", "={l_120}*{t}", "—", "λ · t = 120 / 640 = 0.1875", "0.0000"),
            ("Вероятность безотказной работы", "P_120", "EXP(-λ · t)", "=EXP(-{lt})", "—", "P(120) = exp(-0.1875) ≈ 0.8290 (в слайде округлено до 0.83)", "0.0000"),
            ("Частота отказов регулятора", "f_120", "λ · P(t)", "={l_120}*{P_120}", "1/ч", "f(120) = 0.001563 · 0.8290 ≈ 0.001295 ч⁻¹ (1.29·10⁻³ ч⁻¹)", "0.000000")
        ],
        "При t = 120 ч: интенсивность отказов λ = 1,56·10⁻³ 1/ч; вероятность безотказной работы P(120) = 0,8290 ≈ 0,83; частота отказов f(120) = 1,29·10⁻³ 1/ч (полное совпадение со слайдом 8).",
        cur_r_typ,
        "typ1"
    )

    # --- ТИПОВАЯ ЗАДАЧА 2 ---
    cur_r_typ = render_dynamic_task(
        ws_typ,
        "ЗАДАЧА № 2. Надежность системы из 1000 и 100 элементов при P(t) = 0,99",
        "Система состоит из 1000 элементов с высокой вероятностью безотказной работы P(t) = 0,99. Определить надежность всей системы. Определить надежность этой системы, если число элементов – 100?",
        [
            ("Вероятность безотказной работы элемента", "P_el", 0.99, "—", "P(t) = 0.99 для каждого элемента"),
            ("Число элементов в первой конфигурации", "N1", 1000, "шт.", "N1 = 1000 элементов"),
            ("Число элементов во второй конфигурации", "N2", 100, "шт.", "N2 = 100 элементов")
        ],
        [
            ("Надежность системы при N = 1000", "Pc_1000", "P_el ^ N1", "={P_el}^{N1}", "—", "Pc = 0.99¹⁰⁰⁰ ≈ 0.00004317 (0.432·10⁻⁴ по слайду)", "0.000000"),
            ("Надежность системы при N = 100", "Pc_100", "P_el ^ N2", "={P_el}^{N2}", "—", "Pc = 0.99¹⁰⁰ ≈ 0.36603 (0.366 по слайду)", "0.0000")
        ],
        "При N = 1000 элементов: Pc(t) = 0,432·10⁻⁴ (система крайне ненадежна из-за большого числа элементов); при N = 100 элементов: Pc(t) = 0,366 (36,6%). Полное совпадение со слайдом 8.",
        cur_r_typ,
        "typ2"
    )

    # --- ТИПОВАЯ ЗАДАЧА 3 ---
    cur_r_typ = render_dynamic_task(
        ws_typ,
        "ЗАДАЧА № 3. Система управления из 6000 элементов (T = 50 ч)",
        "Система управления состоит из 6000 элементов, соединенных последовательно, средняя интенсивность отказов которых λср(t) = 0,2∙10⁻⁶ ч⁻¹. Необходимо определить вероятность безотказной работы Pc(t) в течение Т = 50 ч и среднее время безотказной работы системы mtc.",
        [
            ("Число элементов", "n", 6000, "шт.", "Последовательное соединение"),
            ("Интенсивность отказов элемента", "l_elem", 0.0000002, "1/ч", "λср = 0.2·10⁻⁶ ч⁻¹"),
            ("Время работы", "T", 50, "ч", "T = 50 ч")
        ],
        [
            ("Интенсивность отказов системы", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 6000 · 0.2·10⁻⁶ = 0.0012 ч⁻¹ (1.2·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель экспоненты (λc · T)", "lcT", "λc · T", "={lc}*{T}", "—", "λc · T = 0.0012 · 50 = 0.0600 (на слайде опечатка 0.006)", "0.0000"),
            ("Точная вероятность безотказной работы", "Pc_50", "EXP(-λc · T)", "=EXP(-{lcT})", "—", "Pc(50) = exp(-0.06) ≈ 0.9418 (94.18%)", "0.0000"),
            ("Значение по тексту слайда (с опечаткой)", "Pc_slide", "EXP(-0.006)", "=EXP(-0.006)", "—", "В слайде 9 ошибочно записано e^(-0.006) = 0.994", "0.0000"),
            ("Среднее время безотказной работы системы", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.0012 = 833.33 ч (по слайду 833.3 ч)", "0.00")
        ],
        "Результаты: интенсивность отказов λc = 1,2·10⁻³ ч⁻¹; средняя наработка mtc = 833,33 ч; точная вероятность Pc(50) = 0,9418 (в формуле слайда 9 опечатка в степени, приведшая к 0,994).",
        cur_r_typ,
        "typ3"
    )

    # --- ТИПОВАЯ ЗАДАЧА 4 ---
    cur_r_typ = render_dynamic_task(
        ws_typ,
        "ЗАДАЧА № 4. Система из 11000 элементов (t = 80 ч)",
        "Система состоит из 11000 элементов, средняя интенсивность отказов которых λср = 0,12∙10⁻⁶ ч⁻¹. Требуется определить Pc(t), qc(t), fc(t), mtc для t = 80 ч.",
        [
            ("Число элементов", "n", 11000, "шт.", "Количество элементов системы"),
            ("Интенсивность отказов элемента", "l_elem", 0.00000012, "1/ч", "λср = 0.12·10⁻⁶ ч⁻¹"),
            ("Время наработки", "t", 80, "ч", "Контрольная точка t = 80 ч")
        ],
        [
            ("Интенсивность отказов системы", "lc", "n · λср", "={n}*{l_elem}", "1/ч", "λc = 11000 · 0.12·10⁻⁶ = 0.00132 ч⁻¹ (1.32·10⁻³ ч⁻¹)", "0.000000"),
            ("Показатель степени (λc · t)", "lct", "λc · t", "={lc}*{t}", "—", "λc · t = 0.00132 · 80 = 0.1056", "0.0000"),
            ("Вероятность безотказной работы (точная)", "Pc_80", "EXP(-λc · t)", "=EXP(-{lct})", "—", "Pc(80) = exp(-0.1056) ≈ 0.8998 (на слайде указано ≈ 0.82)", "0.0000"),
            ("Вероятность отказа (точная)", "qc_80", "1 - Pc(t)", "=1-{Pc_80}", "—", "qc(80) = 1 - 0.8998 ≈ 0.1002 (на слайде ≈ 0.18)", "0.0000"),
            ("Частота отказов (точная)", "fc_80", "λc · Pc(t)", "={lc}*{Pc_80}", "1/ч", "fc(80) = 0.00132 · 0.8998 ≈ 0.001188 ч⁻¹", "0.000000"),
            ("Среднее время безотказной работы системы", "mtc", "1 / λc", "=1/{lc}", "ч", "mtc = 1 / 0.00132 ≈ 757.58 ч (на слайде ≈ 758 ч)", "0.00")
        ],
        "Результаты: интенсивность отказов λc = 1,32·10⁻³ ч⁻¹; средняя наработка mtc = 757,58 ч ≈ 758 ч; точная Pc(80) = 0,8998; qc(80) = 0,1002; fc(80) = 1,19·10⁻³ ч⁻¹.",
        cur_r_typ,
        "typ4"
    )

    # --- ТИПОВАЯ ЗАДАЧА 5 ---
    cur_r_typ = render_dynamic_task(
        ws_typ,
        "ЗАДАЧА № 5. Расчет характеристик элемента при λ = 2,5∙10⁻⁵ ч⁻¹ (t = 1000 ч)",
        "Время работы элемента до отказа подчинено экспоненциальному закону распределения с параметром λ = 2,5∙10⁻⁵ ч⁻¹. Требуется вычислить количественные характеристики надежности элемента: p(t), q(t), f(t), mt для t = 1000 ч.",
        [
            ("Интенсивность отказов элемента", "l_val", 0.000025, "1/ч", "λ = 2.5·10⁻⁵ ч⁻¹"),
            ("Время наработки", "t", 1000, "ч", "t = 1000 ч")
        ],
        [
            ("Показатель степени (λ · t)", "lt", "λ · t", "={l_val}*{t}", "—", "λ · t = 2.5·10⁻⁵ · 1000 = 0.0250", "0.0000"),
            ("Вероятность безотказной работы", "p_1000", "EXP(-λ · t)", "=EXP(-{lt})", "—", "p(1000) = exp(-0.025) ≈ 0.9753 (по слайду 0.9753)", "0.0000"),
            ("Вероятность отказа элемента", "q_1000", "1 - p(t)", "=1-{p_1000}", "—", "q(1000) = 1 - 0.9753 = 0.0247 (по слайду 0.0247)", "0.0000"),
            ("Частота отказов элемента", "f_1000", "λ · p(t)", "={l_val}*{p_1000}", "1/ч", "f(1000) = 2.5·10⁻⁵ · 0.9753 ≈ 2.439·10⁻⁵ ч⁻¹", "0.00E+00"),
            ("Среднее время безотказной работы", "mt", "1 / λ", "=1/{l_val}", "ч", "mt = 1 / (2.5·10⁻⁵) = 40 000 ч (по слайду 40000 ч)", "#,##0")
        ],
        "При t = 1000 ч: p(1000) = 0,9753; q(1000) = 0,0247; f(1000) = 2,439·10⁻⁵ 1/ч; mt = 40 000 ч. Полное 100% совпадение со слайдами 9–10.",
        cur_r_typ,
        "typ5"
    )


    # =========================================================================
    # CREATE SHEET: СВОДКА И ОТВЕТЫ (INSERT AT INDEX 0)
    # =========================================================================
    ws_sum = wb.create_sheet(title="Сводка и Ответы", index=0)
    ws_sum.views.sheetView[0].showGridLines = True
    
    make_title_banner(
        ws_sum,
        "ПРАКТИЧЕСКАЯ РАБОТА № 2 — СВОДНЫЕ РЕЗУЛЬТАТЫ",
        "Аналитическое определение количественных характеристик надежности изделия (экспоненциальный закон надежности)",
        f"Дисциплина: Надежность автоматизированных систем (НАС)  |  Преподаватель: Ризванов  |  Выполнил: {student_name}",
        max_col=8
    )
    
    ws_sum.cell(row=6, column=2, value="1. РЕЗУЛЬТАТЫ РЕШЕНИЯ ЗАДАЧ ДЛЯ САМОСТОЯТЕЛЬНОГО РЕШЕНИЯ (№ 6 – 12)").font = font_sec_header
    ws_sum.row_dimensions[6].height = 25
    
    sum_headers = [
        "№ задачи", "Краткое условие задачи", "Определяемые параметры", 
        "Расчетная формула (математическая)", "Результат вычислений", "Ед. изм.", "Ссылка на лист решения", "Статус"
    ]
    
    row = 7
    for c_idx, h in enumerate(sum_headers, start=2):
        cell = ws_sum.cell(row=row, column=c_idx, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_tbl_header
        cell.alignment = align_center_wrap
        cell.border = box_cell_border
    ws_sum.row_dimensions[row].height = 28
    
    summary_data = [
        ("Задача 6", "L1=460 ч, t=800 ч. Найти интенсивность отказов λ", "λ(800)", "λ = 1 / L1", f"={RESULTS_REF['t6_l_800']}", "1/ч", RESULTS_REF['t6_start'], "Решено", "0.000000"),
        ("Задача 6 (2)", "Та же задача, вероятность безотказной работы P(800)", "P(800)", "P(t) = exp(-λ · t)", f"={RESULTS_REF['t6_P_800']}", "—", RESULTS_REF['t6_start'], "Решено", "0.0000"),
        ("Задача 6 (3)", "Та же задача, частота отказов f(800)", "f(800)", "f(t) = λ · P(t)", f"={RESULTS_REF['t6_f_800']}", "1/ч", RESULTS_REF['t6_start'], "Решено", "0.000000"),
        ("Задача 7", "n=6000 элементов, λср=0.2·10⁻⁶ ч⁻¹, T=50 ч. Найти Pc(50)", "Pc(50)", "Pc(T) = exp(-n · λср · T)", f"={RESULTS_REF['t7_Pc_50']}", "—", RESULTS_REF['t7_start'], "Решено", "0.0000"),
        ("Задача 7 (2)", "Та же задача, среднее время безотказной работы mtc", "mtc", "mtc = 1 / (n · λср)", f"={RESULTS_REF['t7_mtc']}", "ч", RESULTS_REF['t7_start'], "Решено", "0.00"),
        ("Задача 8", "n=5000 элементов, λср=0.3·10⁻⁶ ч⁻¹, T=120 ч. Найти Pc(120)", "Pc(120)", "Pc(T) = exp(-n · λср · T)", f"={RESULTS_REF['t8_Pc_120']}", "—", RESULTS_REF['t8_start'], "Решено", "0.0000"),
        ("Задача 8 (2)", "Та же задача, среднее время безотказной работы mtc", "mtc", "mtc = 1 / (n · λср)", f"={RESULTS_REF['t8_mtc']}", "ч", RESULTS_REF['t8_start'], "Решено", "0.00"),
        ("Задача 9", "ЭВМ из n=200000 элементов, λср=0.3·10⁻⁶ ч⁻¹, сутки (t=24 ч)", "Pc(24)", "Pc(24) = exp(-n · λср · 24)", f"={RESULTS_REF['t9_Pc_24']}", "—", RESULTS_REF['t9_start'], "Решено", "0.0000"),
        ("Задача 9 (2)", "Та же задача, среднее время безотказной работы ЭВМ", "mtc", "mtc = 1 / (n · λср)", f"={RESULTS_REF['t9_mtc']}", "ч", RESULTS_REF['t9_start'], "Решено", "0.00"),
        ("Задача 10", "n=12600, λср=0.32·10⁻⁶ ч⁻¹, t=50 ч. Вероятность Pc(50)", "Pc(50)", "Pc(t) = exp(-λc · t)", f"={RESULTS_REF['t10_Pc_50']}", "—", RESULTS_REF['t10_start'], "Решено", "0.0000"),
        ("Задача 10 (2)", "Та же задача, вероятность отказа системы qc(50)", "qc(50)", "qc(t) = 1 - Pc(t)", f"={RESULTS_REF['t10_qc_50']}", "—", RESULTS_REF['t10_start'], "Решено", "0.0000"),
        ("Задача 10 (3)", "Та же задача, частота отказов fc(50)", "fc(50)", "fc(t) = λc · Pc(t)", f"={RESULTS_REF['t10_fc_50']}", "1/ч", RESULTS_REF['t10_start'], "Решено", "0.000000"),
        ("Задача 10 (4)", "Та же задача, средняя наработка системы mtc", "mtc", "mtc = 1 / λc", f"={RESULTS_REF['t10_mtc']}", "ч", RESULTS_REF['t10_start'], "Решено", "0.00"),
        ("Задача 11", "Прибор с λ=0.1·10⁻⁶ ч⁻¹, t=180 ч. Вероятность p(180)", "p(180)", "p(t) = exp(-λ · t)", f"={RESULTS_REF['t11_p_180']}", "—", RESULTS_REF['t11_start'], "Решено", "0.000000"),
        ("Задача 11 (2)", "Та же задача, вероятность отказа прибора q(180)", "q(180)", "q(t) = 1 - p(t)", f"={RESULTS_REF['t11_q_180']}", "—", RESULTS_REF['t11_start'], "Решено", "0.000000"),
        ("Задача 11 (3)", "Та же задача, частота отказов прибора f(180)", "f(180)", "f(t) = λ · p(t)", f"={RESULTS_REF['t11_f_180']}", "1/ч", RESULTS_REF['t11_start'], "Решено", "0.00E+00"),
        ("Задача 11 (4)", "Та же задача, среднее время безотказной работы mt", "mt", "mt = 1 / λ", f"={RESULTS_REF['t11_mt']}", "ч", RESULTS_REF['t11_start'], "Решено", "#,##0"),
        ("Задача 12", "Система из 5 приборов (mti={83, 220, 280, 400, 700} ч)", "λc", "λc = Σ (1 / mti)", f"={RESULTS_REF['t12_lc']}", "1/ч", RESULTS_REF['t12_start'], "Решено", "0.000000"),
        ("Задача 12 (2)", "Та же задача, среднее время безотказной работы mtc", "mtc", "mtc = 1 / λc", f"={RESULTS_REF['t12_mtc']}", "ч", RESULTS_REF['t12_start'], "Решено", "0.00")
    ]
    
    for item in summary_data:
        row += 1
        ws_sum.row_dimensions[row].height = 21
        is_even = (row % 2 == 0)
        curr_fill = fill_zebra if is_even else fill_white
        
        c_num = ws_sum.cell(row=row, column=2, value=item[0])
        c_num.font = font_bold; c_num.alignment = align_center; c_num.fill = curr_fill; c_num.border = box_cell_border
        
        c_cond = ws_sum.cell(row=row, column=3, value=item[1])
        c_cond.font = font_regular; c_cond.alignment = align_left; c_cond.fill = curr_fill; c_cond.border = box_cell_border
        
        c_param = ws_sum.cell(row=row, column=4, value=item[2])
        c_param.font = font_bold; c_param.alignment = align_center; c_param.fill = curr_fill; c_param.border = box_cell_border
        
        c_form = ws_sum.cell(row=row, column=5, value=item[3])
        c_form.font = font_math; c_form.alignment = align_left; c_form.fill = curr_fill; c_form.border = box_cell_border
        
        c_val = ws_sum.cell(row=row, column=6, value=item[4])
        c_val.font = font_result; c_val.alignment = align_right; c_val.fill = fill_result; c_val.border = box_result_border
        c_val.number_format = item[8]
        
        c_unit = ws_sum.cell(row=row, column=7, value=item[5])
        c_unit.font = font_regular; c_unit.alignment = align_center; c_unit.fill = curr_fill; c_unit.border = box_cell_border
        
        c_link = ws_sum.cell(row=row, column=8, value="Перейти к решению")
        c_link.hyperlink = f"#{item[6]}"
        c_link.font = Font(name=FONT_NAME, size=9, underline='single', color='1F4E78')
        c_link.alignment = align_center; c_link.fill = curr_fill; c_link.border = box_cell_border
        
        c_stat = ws_sum.cell(row=row, column=9, value=item[7])
        c_stat.font = font_status_ok; c_stat.alignment = align_center; c_stat.fill = curr_fill; c_stat.border = box_cell_border

    # Table 2: Typical Examples Summary
    row += 3
    ws_sum.cell(row=row, column=2, value="2. РЕЗУЛЬТАТЫ РЕШЕНИЯ ТИПОВЫХ ЗАДАЧ ИЗ ПРЕЗЕНТАЦИИ (№ 1 – 5)").font = font_sec_header
    ws_sum.row_dimensions[row].height = 25
    
    row += 1
    typ_headers = [
        "№ задачи", "Тема / Объект расчета", "Ключевые исходные данные",
        "Искомый параметр", "Результат вычислений", "Ед. изм.", "Сверка со слайдом", "Ссылка"
    ]
    for c_idx, h in enumerate(typ_headers, start=2):
        cell = ws_sum.cell(row=row, column=c_idx, value=h)
        cell.font = font_tbl_header; cell.fill = fill_tbl_header; cell.alignment = align_center_wrap; cell.border = box_cell_border
    ws_sum.row_dimensions[row].height = 26
    
    typ_summary_data = [
        ("Задача 1", "Автоматический регулятор", "L1=640 ч, t=120 ч", "P(120), λ, f(120)", f"={RESULTS_REF['typ1_P_120']}", "—", "P=0.83; λ=1.56·10⁻³; f=1.29·10⁻³ (Совпадает)", RESULTS_REF['typ1_start']),
        ("Задача 2", "Система из N одинаковых элементов", "P=0.99, N1=1000, N2=100", "Pc(1000), Pc(100)", f"={RESULTS_REF['typ2_Pc_100']}", "—", "Pc(1000)=0.432·10⁻⁴; Pc(100)=0.366 (Совпадает)", RESULTS_REF['typ2_start']),
        ("Задача 3", "Система управления", "n=6000, λср=0.2·10⁻⁶, T=50 ч", "Pc(50), mtc", f"={RESULTS_REF['typ3_mtc']}", "ч", "mtc=833.3 ч; Pc(50)=0.9418 (в слайде опечатка 0.994)", RESULTS_REF['typ3_start']),
        ("Задача 4", "Система из 11000 элементов", "n=11000, λср=0.12·10⁻⁶, t=80 ч", "Pc(80), mtc, fc(80)", f"={RESULTS_REF['typ4_mtc']}", "ч", "mtc=758 ч; Pc=0.8998 (на слайде ≈0.82)", RESULTS_REF['typ4_start']),
        ("Задача 5", "Элемент (экспоненциальный закон)", "λ=2.5·10⁻⁵ ч⁻¹, t=1000 ч", "p(1000), q, f, mt", f"={RESULTS_REF['typ5_p_1000']}", "—", "p=0.9753; q=0.0247; f=2.439·10⁻⁵; mt=40000 (Совпадает)", RESULTS_REF['typ5_start']),
    ]
    
    for item in typ_summary_data:
        row += 1
        ws_sum.row_dimensions[row].height = 21
        is_even = (row % 2 == 0)
        curr_fill = fill_zebra if is_even else fill_white
        
        c_num = ws_sum.cell(row=row, column=2, value=item[0])
        c_num.font = font_bold; c_num.alignment = align_center; c_num.fill = curr_fill; c_num.border = box_cell_border
        
        c_theme = ws_sum.cell(row=row, column=3, value=item[1])
        c_theme.font = font_regular; c_theme.alignment = align_left; c_theme.fill = curr_fill; c_theme.border = box_cell_border
        
        c_inp = ws_sum.cell(row=row, column=4, value=item[2])
        c_inp.font = font_regular; c_inp.alignment = align_left; c_inp.fill = curr_fill; c_inp.border = box_cell_border
        
        c_p = ws_sum.cell(row=row, column=5, value=item[3])
        c_p.font = font_bold; c_p.alignment = align_center; c_p.fill = curr_fill; c_p.border = box_cell_border
        
        c_res = ws_sum.cell(row=row, column=6, value=item[4])
        c_res.font = font_result; c_res.alignment = align_right; c_res.fill = fill_result; c_res.border = box_result_border
        c_res.number_format = '0.0000'
        
        c_u = ws_sum.cell(row=row, column=7, value=item[5])
        c_u.font = font_regular; c_u.alignment = align_center; c_u.fill = curr_fill; c_u.border = box_cell_border
        
        c_chk = ws_sum.cell(row=row, column=8, value=item[6])
        c_chk.font = font_status_ok; c_chk.alignment = align_left; c_chk.fill = curr_fill; c_chk.border = box_cell_border
        
        c_l = ws_sum.cell(row=row, column=9, value="Смотреть")
        c_l.hyperlink = f"#{item[7]}"
        c_l.font = Font(name=FONT_NAME, size=9, underline='single', color='1F4E78')
        c_l.alignment = align_center; c_l.fill = curr_fill; c_l.border = box_cell_border


    # =========================================================================
    # CREATE SHEET: ТЕОРИЯ И ФОРМУЛЫ
    # =========================================================================
    ws_th = wb.create_sheet(title="Теория и формулы")
    ws_th.views.sheetView[0].showGridLines = True
    
    make_title_banner(
        ws_th,
        "ТЕОРЕТИЧЕСКИЙ СПРАВОЧНИК — АНАЛИТИЧЕСКАЯ НАДЕЖНОСТЬ",
        "Аналитическое определение количественных характеристик надежности изделия (ГОСТ 27.002 и тема 2)",
        "Определения, обозначения, математические формулы и физический смысл",
        max_col=7
    )
    
    ws_th.cell(row=6, column=2, value="ОСНОВНЫЕ АНАЛИТИЧЕСКИЕ ХАРАКТЕРИСТИКИ НАДЕЖНОСТИ").font = font_sec_header
    ws_th.row_dimensions[6].height = 25
    
    th_headers = ["Показатель надежности", "Обозначение", "Общая формула", "Формула при эксп. законе", "Ед. изм.", "Физический смысл и пояснение"]
    r = 7
    for ci, th in enumerate(th_headers, start=2):
        cell = ws_th.cell(row=r, column=ci, value=th)
        cell.font = font_tbl_header; cell.fill = fill_tbl_header; cell.alignment = align_center_wrap; cell.border = box_cell_border
    ws_th.row_dimensions[r].height = 26
    
    theory_rows = [
        ("Вероятность безотказной работы", "p(t)", "p(t) = exp(-∫λ(t)dt) = 1 - ∫f(t)dt", "p(t) = e^(-λ·t)", "—", "Вероятность того, что в пределах заданной наработки t отказ объекта не возникнет"),
        ("Вероятность отказа изделия", "q(t)", "q(t) = 1 - p(t)", "q(t) = 1 - e^(-λ·t)", "—", "Вероятность того, что объект откажет хотя бы один раз за время от 0 до t"),
        ("Частота отказов (плотность)", "f(t)", "f(t) = -dp(t)/dt = dq(t)/dt", "f(t) = λ · e^(-λ·t) = λ · p(t)", "1/ч", "Плотность вероятности возникновения отказа невосстанавливаемого объекта в момент t"),
        ("Интенсивность отказов", "λ(t)", "λ(t) = f(t) / p(t) = -(1/p)·dp/dt", "λ(t) = λ = const", "1/ч", "Условная плотность вероятности отказа объекта в момент t при условии безотказной работы до t"),
        ("Средняя наработка до отказа", "mt", "mt = ∫ p(t) dt  (от 0 до ∞)", "mt = 1 / λ", "ч", "Математическое ожидание времени безотказной работы объекта")
    ]
    
    for trow in theory_rows:
        r += 1
        ws_th.row_dimensions[r].height = 28
        curr_fill = fill_zebra if r % 2 == 0 else fill_white
        
        c1 = ws_th.cell(row=r, column=2, value=trow[0]); c1.font = font_bold; c1.alignment = align_left; c1.fill = curr_fill; c1.border = box_cell_border
        c2 = ws_th.cell(row=r, column=3, value=trow[1]); c2.font = font_bold; c2.alignment = align_center; c2.fill = curr_fill; c2.border = box_cell_border
        c3 = ws_th.cell(row=r, column=4, value=trow[2]); c3.font = font_math; c3.alignment = align_left; c3.fill = curr_fill; c3.border = box_cell_border
        c4 = ws_th.cell(row=r, column=5, value=trow[3]); c4.font = font_bold; c4.alignment = align_left; c4.fill = curr_fill; c4.border = box_cell_border
        c5 = ws_th.cell(row=r, column=6, value=trow[4]); c5.font = font_regular; c5.alignment = align_center; c5.fill = curr_fill; c5.border = box_cell_border
        c6 = ws_th.cell(row=r, column=7, value=trow[5]); c6.font = font_italic; c6.alignment = align_left_wrap; c6.fill = curr_fill; c6.border = box_cell_border
        
    r += 2
    ws_th.cell(row=r, column=2, value="ОСОБЕННОСТИ ЭКСПОНЕНЦИАЛЬНОГО ЗАКОНА И СИСТЕМНОЙ НАДЕЖНОСТИ").font = font_sec_header
    ws_th.row_dimensions[r].height = 25
    
    rel_rows = [
        ("Свойство отсутствия последействия (старения):", "P(t + τ | t) = P(τ) = exp(-λ · τ). Объект не изнашивается и не стареет; надежность зависит только от длительности предстоящей работы."),
        ("Постоянство интенсивности отказов:", "λ(t) = λ = const на всем периоде нормальной эксплуатации изделия (участок II кривой интенсивности отказов)."),
        ("Связь средней наработки и интенсивности:", "mt = 1 / λ   <=>   λ = 1 / mt. Если средняя наработка известна, интенсивность отказов находится обратным делением."),
        ("Последовательное соединение n элементов:", "λc = Σ λi (теорема сложения интенсивностей); для одинаковых элементов: λc = n · λср."),
        ("Вероятность безотказной работы системы:", "Pc(t) = Π Pi(t) = exp(-λc · t) = exp(-Σ λi · t); для одинаковых элементов: Pc(t) = [P(t)]ⁿ = exp(-n · λср · t)."),
        ("Средняя наработка последовательной системы:", "mtc = 1 / λc = 1 / (Σ 1 / mti). Средняя наработка системы всегда меньше средней наработки самого ненадежного элемента!")
    ]
    for rel_title, rel_expr in rel_rows:
        r += 1
        ws_th.row_dimensions[r].height = 26
        ws_th.cell(row=r, column=2, value=rel_title).font = font_bold
        ws_th.cell(row=r, column=2).alignment = align_left
        ws_th.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        c_rel = ws_th.cell(row=r, column=3, value=rel_expr)
        c_rel.font = font_regular
        c_rel.alignment = align_left_wrap
        apply_row_style(ws_th, r, 2, 7, fill=fill_condition, border=box_cell_border)

    # Column widths
    column_widths = {
        "Сводка и Ответы": {
            "A": 4, "B": 18, "C": 44, "D": 26, "E": 36, "F": 22, "G": 10, "H": 22, "I": 12
        },
        "Задачи 6-12": {
            "A": 4, "B": 42, "C": 18, "D": 34, "E": 24, "F": 12, "G": 48
        },
        "Задачи 1-5": {
            "A": 4, "B": 42, "C": 18, "D": 34, "E": 24, "F": 12, "G": 48
        },
        "Теория и формулы": {
            "A": 4, "B": 36, "C": 16, "D": 36, "E": 28, "F": 12, "G": 52
        }
    }
    
    for sheet_name, widths in column_widths.items():
        if sheet_name in wb.sheetnames:
            ws_curr = wb[sheet_name]
            for col_letter, width in widths.items():
                ws_curr.column_dimensions[col_letter].width = width

    # Save to disk
    wb.save(filename_out)
    print(f"Successfully generated: {filename_out}")
    return wb

if __name__ == '__main__':
    # 1. Master file for Gainanov
    generate_full_practice_2(
        student_name="Студент Гайнанов Д.И. (группа ЭАС-414С)",
        filename_out="НАС_ПР2_ГайнановДИ.xlsx"
    )
    # 2. Master general reference file
    generate_full_practice_2(
        student_name="Студент",
        filename_out="Практика_2_Решение.xlsx"
    )
