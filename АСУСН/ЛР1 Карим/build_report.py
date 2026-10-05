"""
Скрипт генерации отчета по Лабораторной работе №1
по дисциплине 'Автоматизированные системы специального назначения'
Форматирование: строгое академическое (по ГОСТ и примеру),
обычные таблицы без цветовых заливок,
обычные листинги без цветов,
оригинальные программы и листинги из методических указаний.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Установка внутренних отступов ячейки таблицы"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="000000", sz="4", val="single"):
    """Установка стандартных одинарных границ таблицы черного цвета"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_page_number(run):
    """Вставка динамического поля номера страницы в колонтитул Word"""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)


def add_code_block(doc, code_text):
    """Обычное строгое форматирование листинга без цветов (моноширинный шрифт, тонкая рамка)"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    
    # Тонкая одинарная черная рамка без фоновой заливки (без цветов)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)


def create_report():
    doc = docx.Document()
    
    # Настройка страницы (А4, поля по ГОСТ как в примере)
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Pt(56.7)     # 2.0 см
    sec.bottom_margin = Pt(56.7)  # 2.0 см
    sec.left_margin = Pt(85.05)   # 3.0 см
    sec.right_margin = Pt(42.5)   # 1.5 см

    # Колонтитулы и нумерация страниц со второй страницы
    sec.different_first_page_header_footer = True
    footer = sec.footer
    p_f = footer.paragraphs[0]
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_f = p_f.add_run()
    run_f.font.name = 'Times New Roman'
    run_f.font.size = Pt(11)
    add_page_number(run_f)

    # Базовый стиль шрифта
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0, 0, 0)

    # =============================================================
    # ТИТУЛЬНЫЙ ЛИСТ
    # =============================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Федеральное государственное бюджетное образовательное учреждение\nвысшего образования")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("«Уфимский университет науки и технологий»")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(60)
    r = p.add_run("Кафедра АСУ")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Отчет по лабораторной работе № 1")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("по дисциплине: Автоматизированные системы специального назначения")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(100)
    r = p.add_run("на тему: «Изучение архитектуры и принципов работы программируемого логического контроллера (ПЛК) на основе программного моделирования»")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True

    # Блок автора и проверяющего справа
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Выполнил:")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("студент гр. ЭАС-514С")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Гарифуллин Карим Рамилевич")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Проверил:")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(90)
    r = p.add_run("Антонов В.В.")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Уфа 2026 г.")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    doc.add_page_break()

    # Вспомогательные функции
    def add_sec_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        return p

    def add_sub_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12.5)
        r.font.bold = True
        return p

    def add_body_p(text, indent=35.4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(indent)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_bullet_p(bold_prefix, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Pt(25)
        p.paragraph_format.first_line_indent = Pt(-15)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        r_bullet = p.add_run("• ")
        r_bullet.font.bold = True
        if bold_prefix:
            r_bold = p.add_run(bold_prefix + " ")
            r_bold.font.bold = True
        r_text = p.add_run(text)
        return p

    def add_figure(img_path, caption, width=Inches(6.0)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.first_line_indent = Pt(0)
            doc.add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            p_cap.paragraph_format.first_line_indent = Pt(0)
            r = p_cap.add_run(caption)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.italic = True

    # =============================================================
    # ЦЕЛЬ РАБОТЫ
    # =============================================================
    add_sec_title("Цель работы")
    add_body_p(
        "Изучить архитектуру и принципы функционирования программируемого логического контроллера (ПЛК); "
        "освоить методы описания логики работы дискретных и аналоговых входов/выходов; получить навыки "
        "моделирования поведения ПЛК с использованием общедоступных программных средств."
    )

    # =============================================================
    # ТЕОРЕТИЧЕСКАЯ ЧАСТЬ
    # =============================================================
    add_sec_title("Теоретическая часть")
    add_body_p(
        "ПЛК – промышленное устройство, предназначенное для управления технологическими процессами "
        "в реальном времени. Типовая архитектура включает: центральный процессор (CPU), модули ввода/вывода (I/O), "
        "блок питания, коммуникационные модули и подсистему памяти."
    )
    add_body_p(
        "Цикл работы ПЛК состоит из периодически повторяющихся этапов:"
    )
    add_bullet_p("1.", "Чтение входов (Input Scan);")
    add_bullet_p("2.", "Выполнение программы (Program Execution);")
    add_bullet_p("3.", "Запись выходов (Output Update);")
    add_bullet_p("4.", "Обслуживание коммуникаций (Communication);")
    add_bullet_p("5.", "Самодиагностика (Self-Diagnostics);")
    add_bullet_p("6.", "Возврат к шагу 1.")
    add_body_p(
        "Время цикла – ключевая характеристика ПЛК. Для систем реального времени критично, чтобы цикл "
        "укладывался в строго заданный интервал."
    )
    add_body_p(
        "Дискретные входы воспринимают сигналы типа «включено/выключено» (уровни 24 В DC, 220 В AC). "
        "Дискретные выходы управляют исполнительными механизмами (реле, транзисторы, симисторы). "
        "Аналоговые входы и выходы работают с непрерывными сигналами (стандарты 4–20 мА, 0–10 В) "
        "через преобразование АЦП/ЦАП в целочисленные значения (0–27648)."
    )

    # =============================================================
    # ХОД РАБОТЫ
    # =============================================================
    add_sec_title("Ход работы")
    add_sub_title("Часть 1. Анализ архитектуры ПЛК")
    add_body_p(
        "На основе методических указаний составлена переработанная структурная схема ПЛК "
        "с указанием основных внутренних функциональных компонентов и внешних полевых устройств (рисунок 1)."
    )

    add_figure("plc_architecture.png", "Рисунок 1 – Структурная схема программируемого логического контроллера")

    add_body_p(
        "Временные характеристики этапов цикла работы ПЛК представлены в таблице 1. Суммарное время цикла "
        "составляет порядка 10 мс, что обеспечивает детерминированную реакцию системы управления."
    )

    # Таблица 1: Цикл работы ПЛК (обычная без цветов)
    p_tbl1_cap = doc.add_paragraph()
    p_tbl1_cap.paragraph_format.space_before = Pt(8)
    p_tbl1_cap.paragraph_format.space_after = Pt(2)
    p_tbl1_cap.paragraph_format.first_line_indent = Pt(0)
    r = p_tbl1_cap.add_run("Таблица 1 – Этапы цикла работы ПЛК")
    r.font.bold = True
    r.font.size = Pt(11)

    t1 = doc.add_table(rows=7, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)

    t1_headers = ["Этап", "Действие", "Время выполнения"]
    for j, h in enumerate(t1_headers):
        cell = t1.cell(0, j)
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)

    t1_data = [
        ("1", "Чтение входов", "~1 мс"),
        ("2", "Выполнение программы", "~5 мс"),
        ("3", "Запись выходов", "~1 мс"),
        ("4", "Обслуживание коммуникаций", "~2 мс"),
        ("5", "Самодиагностика", "~1 мс"),
        ("Итого", "Время цикла", "~10 мс")
    ]

    for i, row in enumerate(t1_data, start=1):
        for j, val in enumerate(row):
            cell = t1.cell(i, j)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if j == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(10)
            if i == 6 or j == 0:
                r.font.bold = True

    add_body_p(
        "В таблице 2 приведено сопоставление характеристик контроллеров различных производителей "
        "(Siemens, Allen-Bradley, Schneider Electric, ОВЕН)."
    )

    # Таблица 2: Сравнение ПЛК (обычная без цветов)
    p_tbl2_cap = doc.add_paragraph()
    p_tbl2_cap.paragraph_format.space_before = Pt(8)
    p_tbl2_cap.paragraph_format.space_after = Pt(2)
    p_tbl2_cap.paragraph_format.first_line_indent = Pt(0)
    r = p_tbl2_cap.add_run("Таблица 2 – Сравнение ПЛК различных производителей")
    r.font.bold = True
    r.font.size = Pt(11)

    t2 = doc.add_table(rows=8, cols=5)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)

    t2_headers = ["Параметр", "Siemens S7-1200", "Allen-Bradley CompactLogix", "Schneider Modicon M241", "ОВЕН ПЛК160"]
    for j, h in enumerate(t2_headers):
        cell = t2.cell(0, j)
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)

    t2_data = [
        ("Производитель", "Siemens", "Rockwell Automation", "Schneider Electric", "ОВЕН"),
        ("Языки программирования", "LD, FBD, ST, SFC", "LD, FBD, ST, SFC", "LD, FBD, ST, SFC", "LD, FBD, ST, SFC"),
        ("Дискретные входы", "до 14", "до 32", "до 14", "до 16"),
        ("Дискретные выходы", "до 10", "до 32", "до 10", "до 12"),
        ("Аналоговые входы", "до 2", "до 8", "до 2", "до 8"),
        ("Интерфейсы", "Profinet, Modbus TCP", "EtherNet/IP, Modbus", "Ethernet, Modbus", "Ethernet, Modbus"),
        ("Среда программирования", "TIA Portal", "Studio 5000", "EcoStruxure", "CODESYS")
    ]

    for i, row in enumerate(t2_data, start=1):
        for j, val in enumerate(row):
            cell = t2.cell(i, j)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if j == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True

    # =============================================================
    # ЧАСТЬ 2. МОДЕЛИРОВАНИЕ ДИСКРЕТНОЙ ЛОГИКИ
    # =============================================================
    add_sub_title("Часть 2. Моделирование дискретной логики")
    add_body_p(
        "В среде Microsoft Excel была создана таблица истинности для базовых логических операций "
        "«И», «ИЛИ», «НЕ», «XOR» с использованием встроенных формул (таблица 3)."
    )

    # Таблица 3: Таблица истинности (обычная без цветов)
    p_tbl3_cap = doc.add_paragraph()
    p_tbl3_cap.paragraph_format.space_before = Pt(8)
    p_tbl3_cap.paragraph_format.space_after = Pt(2)
    p_tbl3_cap.paragraph_format.first_line_indent = Pt(0)
    r = p_tbl3_cap.add_run("Таблица 3 – Таблица истинности логических операций")
    r.font.bold = True
    r.font.size = Pt(11)

    t3 = doc.add_table(rows=5, cols=6)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)

    t3_headers = ["A", "B", "A И B", "A ИЛИ B", "НЕ A", "A XOR B"]
    for j, h in enumerate(t3_headers):
        cell = t3.cell(0, j)
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)

    t3_data = [
        ("0", "0", "0", "0", "1", "0"),
        ("0", "1", "0", "1", "1", "1"),
        ("1", "0", "0", "1", "0", "1"),
        ("1", "1", "1", "1", "0", "0")
    ]

    for i, row in enumerate(t3_data, start=1):
        for j, val in enumerate(row):
            cell = t3.cell(i, j)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(10)

    add_body_p(
        "В файле lab1_discrete.py реализована модель дискретной логики ПЛК (класс PLC_Discrete), "
        "включающая управление лампой по кнопке с самоблокировкой, логику «И», «ИЛИ», «НЕ» и счётчик нажатий кнопки. "
        "Ниже представлен листинг программы lab1_discrete.py:"
    )

    with open('lab1_discrete.py', 'r', encoding='utf-8') as f:
        discrete_code = f.read()
    add_code_block(doc, discrete_code)

    add_body_p(
        "Результаты тестирования сценариев работы дискретной логики:"
    )
    add_bullet_p("Сценарий 1 (Нажатие пуска):", "DI0=True, DI1=False, DI2=False, DI3=False -> DQ0=True, DQ1=False, DQ2=True, DQ3=True, Счётчик=1")
    add_bullet_p("Сценарий 2 (Отпускание пуска):", "DI0=False, DI1=False, DI2=False, DI3=False -> DQ0=True, DQ1=False, DQ2=False, DQ3=True, Счётчик=1 (самоблокировка активна)")
    add_bullet_p("Сценарий 3 (Нажатие стопа):", "DI0=False, DI1=True, DI2=False, DI3=False -> DQ0=False, DQ1=False, DQ2=False, DQ3=False, Счётчик=1")
    add_bullet_p("Сценарий 4 (Логика И):", "DI0=True, DI1=False, DI2=True, DI3=False -> DQ0=True, DQ1=True, DQ2=True, DQ3=True, Счётчик=2")
    add_bullet_p("Сценарий 5 (Логика ИЛИ):", "DI0=True, DI1=False, DI2=False, DI3=True -> DQ0=True, DQ1=False, DQ2=True, DQ3=True, Счётчик=2")

    add_body_p(
        "На основе полученных результатов построены временные диаграммы работы дискретных сигналов (рисунок 2)."
    )
    add_figure("timing_diagram.png", "Рисунок 2 – Временные диаграммы работы дискретной логики")

    # =============================================================
    # ЧАСТЬ 3. МОДЕЛИРОВАНИЕ АНАЛОГОВОГО СИГНАЛА
    # =============================================================
    add_sub_title("Часть 3. Моделирование аналогового сигнала")
    add_body_p(
        "В среде Microsoft Excel реализовано масштабирование аналогового сигнала: "
        "преобразование «сырого» значения АЦП (0–27648) в физическую величину (0–100%) по формуле =A/27648*100 (таблица 4)."
    )

    # Таблица 4: Масштабирование (обычная без цветов)
    p_tbl4_cap = doc.add_paragraph()
    p_tbl4_cap.paragraph_format.space_before = Pt(8)
    p_tbl4_cap.paragraph_format.space_after = Pt(2)
    p_tbl4_cap.paragraph_format.first_line_indent = Pt(0)
    r = p_tbl4_cap.add_run("Таблица 4 – Масштабирование аналогового сигнала в Excel")
    r.font.bold = True
    r.font.size = Pt(11)

    t4 = doc.add_table(rows=6, cols=3)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t4)

    t4_headers = ["A (сырое)", "B (формула)", "C (результат)"]
    for j, h in enumerate(t4_headers):
        cell = t4.cell(0, j)
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)

    t4_data = [
        ("0", "=A2/27648*100", "0.0"),
        ("6912", "=A3/27648*100", "25.0"),
        ("13824", "=A4/27648*100", "50.0"),
        ("20736", "=A5/27648*100", "75.0"),
        ("27648", "=A6/27648*100", "100.0")
    ]

    for i, row in enumerate(t4_data, start=1):
        for j, val in enumerate(row):
            cell = t4.cell(i, j)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if j == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(10)

    add_body_p(
        "В файле lab1_analog.py реализован класс PLC_Analog, выполняющий прямое масштабирование входа "
        "scale_input и обратное преобразование scale_output для аналогового выхода. "
        "Ниже представлен листинг программы lab1_analog.py:"
    )

    with open('lab1_analog.py', 'r', encoding='utf-8') as f:
        analog_code = f.read()
    add_code_block(doc, analog_code)

    add_body_p(
        "Результаты работы программы lab1_analog.py при тестировании контрольных точек:"
    )
    add_bullet_p("Сырое: 0", "Физическое: 0.00 % | Обратно (ЦАП): 0")
    add_bullet_p("Сырое: 6912", "Физическое: 25.00 % | Обратно (ЦАП): 6912")
    add_bullet_p("Сырое: 13824", "Физическое: 50.00 % | Обратно (ЦАП): 13824")
    add_bullet_p("Сырое: 20736", "Физическое: 75.00 % | Обратно (ЦАП): 20736")
    add_bullet_p("Сырое: 27648", "Физическое: 100.00 % | Обратно (ЦАП): 27648")

    add_body_p(
        "Построенный график зависимости выходного физического значения от входного сырого кода АЦП "
        "представлен на рисунке 3."
    )
    add_figure("analog_scaling.png", "Рисунок 3 – Характеристика масштабирования аналогового сигнала")

    # =============================================================
    # ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ
    # =============================================================
    add_sec_title("Ответы на контрольные вопросы")

    questions = [
        ("1. Какие компоненты входят в состав ПЛК?",
         "В состав ПЛК входят: центральный процессор (CPU), подсистема памяти (ROM, RAM, NVRAM), "
         "модули дискретного ввода (DI) и вывода (DQ), модули аналогового ввода (AI) и вывода (AQ), "
         "блок питания (PSU), коммуникационные модули (Ethernet, RS-485) и внутренняя системная шина."),

        ("2. Опишите цикл работы ПЛК.",
         "Цикл работы ПЛК включает последовательные этапы: 1) Чтение состояния физических входов в область "
         "памяти образа процесса (PII); 2) Выполнение пользовательской программы управления; "
         "3) Запись вычисленных значений из образа процесса выходов (PIQ) в физические выходы; "
         "4) Обслуживание коммуникационных запросов и обмен данными по сети; "
         "5) Выполнение внутренней самодиагностики и сброс сторожевого таймера (WDT); 6) Возврат к началу цикла."),

        ("3. Чем отличается дискретный сигнал от аналогового?",
         "Дискретный сигнал может принимать только два логических состояния: «0» (выключено/низкий уровень) "
         "или «1» (включено/высокий уровень, например 24 В). Аналоговый сигнал является непрерывной величиной "
         "(ток 4–20 мА или напряжение 0–10 В), которая может принимать любое значение в заданном диапазоне "
         "и пропорциональна измеряемому технологическому параметру (давлению, температуре, уровню)."),

        ("4. Как масштабируется аналоговый сигнал 4–20 мА?",
         "Аналоговый сигнал оцифровывается модулем АЦП в целочисленное сырое значение (например, 0–27648). "
         "Масштабирование в физическую величину выполняется по линейной формуле интерполяции: "
         "Scaled = MinScale + (Raw - MinRaw) / (MaxRaw - MinRaw) * (MaxScale - MinScale). "
         "Значению 4 мА соответствует код 0 (начало шкалы), а значению 20 мА соответствует код 27648 (конец шкалы)."),

        ("5. Какие логические операции используются в дискретной логике?",
         "В дискретной логике используются базовые булевы операции: логическое «И» (AND, активно, если все входы активны), "
         "логическое «ИЛИ» (OR, активно, если активен хотя бы один вход), логическое «НЕ» (NOT, инвертирует состояние входа), "
         "а также «Исключающее ИЛИ» (XOR, активно, если состояния входов различаются)."),

        ("6. Как реализовать самоблокировку в логике управления?",
         "Самоблокировка реализуется включением контакта выходного реле параллельно кнопке «Пуск» и последовательно "
         "с размыкающей кнопкой «Стоп». В программной логике: выход DQ устанавливается в True при нажатии Пуск и сбрасывается "
         "в False при нажатии Стоп. При отпускании кнопки Пуск сигнал удерживается за счет собственного состояния выхода DQ."),

        ("7. Какие средства можно использовать для моделирования ПЛК на обычном ПК?",
         "Для моделирования ПЛК на ПК можно применять: язык Python (моделирование логических алгоритмов, библиотека matplotlib "
         "для временных диаграмм, tkinter для мнемосхем); табличные процессоры Excel / Calc (таблицы истинности, формулы AND/OR/NOT, "
         "масштабирование сигналов); среды моделирования цифровых схем (Logisim, Digital Works); специализированные эмуляторы ПЛК (Siemens PLCSIM, CODESYS)."),

        ("8. Почему время цикла критично для систем реального времени?",
         "Время цикла определяет быстродействие системы управления. Если время цикла слишком велико или нестабильно, "
         "контроллер не успеет своевременно среагировать на быстропротекающие аварийные события технологического процесса "
         "(например, резкий скачок давления, перегрузка), что может привести к выходу оборудования из строя или аварии.")
    ]

    for q, ans in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.first_line_indent = Pt(0)
        p_q.paragraph_format.keep_with_next = True
        r = p_q.add_run(q)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True

        p_a = doc.add_paragraph()
        p_a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_a.paragraph_format.first_line_indent = Pt(35.4)
        p_a.paragraph_format.line_spacing = 1.15
        p_a.paragraph_format.space_after = Pt(4)
        r = p_a.add_run(ans)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)

    # =============================================================
    # ВЫВОДЫ
    # =============================================================
    add_sec_title("Вывод")
    add_body_p(
        "В ходе выполнения лабораторной работы была изучена архитектура и принципы функционирования "
        "программируемых логических контроллеров (ПЛК). Были решены следующие задачи:"
    )
    add_bullet_p("1.", "Составлена структурная схема ПЛК, отражающая взаимодействие центрального процессора, подсистемы памяти, "
                 "блока питания, коммуникационных модулей и модулей ввода/вывода (DI, DQ, AI, AQ) с внешними датчиками и исполнительными механизмами.")
    add_bullet_p("2.", "Проанализирован цикл работы ПЛК, состоящий из фаз опроса входов, выполнения программы, обновления выходов, "
                 "коммуникации и самодиагностики с суммарным временем цикла порядка 10 мс.")
    add_bullet_p("3.", "Проведено сравнительное сопоставление промышленных контроллеров ведущих производителей (Siemens, Allen-Bradley, Schneider Electric, ОВЕН).")
    add_bullet_p("4.", "В среде Excel реализована таблица истинности базовых операций дискретной логики («И», «ИЛИ», «НЕ», «XOR»), "
                 "а на языке Python в файле lab1_discrete.py разработана программная модель дискретной логики с самоблокировкой и счетчиком, "
                 "построены временные диаграммы работы сигналов.")
    add_bullet_p("5.", "В среде Excel и на языке Python в файле lab1_analog.py реализовано прямое и обратное масштабирование аналогового сигнала "
                 "(преобразование сырого значения АЦП 0–27648 в физическую величину 0–100%), построен график характеристики масштабирования.")

    output_filename = "Отчет_ЭАС-514С_АССН_Гарифуллин_ЛР1.docx"
    doc.save(output_filename)
    print(f"Отчет успешно сохранен в '{output_filename}'!")


if __name__ == '__main__':
    create_report()
